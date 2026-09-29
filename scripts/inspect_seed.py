"""
Inspect the seed dataset Excel file.
Reports: sheet names, columns, row counts, sample data, and data quality issues.
"""
import sys
import json
import openpyxl

XLSX_PATH = r"c:\Nisha\PS26108_Seed_Dataset_v1.xlsx"

def inspect():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    report = {"sheets": {}}
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        
        if not rows:
            report["sheets"][sheet_name] = {"status": "EMPTY", "rows": 0}
            continue
        
        headers = [str(h).strip() if h is not None else f"__col_{i}" for i, h in enumerate(rows[0])]
        data_rows = rows[1:]
        
        # Column analysis
        col_info = {}
        for ci, col_name in enumerate(headers):
            values = [r[ci] for r in data_rows if ci < len(r)]
            non_null = [v for v in values if v is not None and str(v).strip() != ""]
            unique = set(str(v) for v in non_null)
            col_info[col_name] = {
                "total": len(values),
                "non_null": len(non_null),
                "null": len(values) - len(non_null),
                "unique": len(unique),
                "sample_values": list(unique)[:5]
            }
        
        # First 3 rows as dicts
        sample_rows = []
        for r in data_rows[:3]:
            row_dict = {}
            for ci, col_name in enumerate(headers):
                val = r[ci] if ci < len(r) else None
                row_dict[col_name] = str(val) if val is not None else None
            sample_rows.append(row_dict)
        
        report["sheets"][sheet_name] = {
            "headers": headers,
            "row_count": len(data_rows),
            "columns": col_info,
            "sample_rows": sample_rows,
        }
    
    wb.close()
    print(json.dumps(report, indent=2, default=str))

if __name__ == "__main__":
    inspect()
